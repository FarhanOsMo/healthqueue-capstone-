"""Build reports/dora-analysis.xlsx from the synthetic 90-day deployment dataset.

All metrics are live Excel formulas over the Raw_Data sheet (named ranges such
as LeadTimeHours make them readable). The Checks sheet compares each formula
with an independent Python calculation; every row must show PASS.

Input:  data/dora_deployments.csv   (python scripts/generate_dora_data.py)
Usage:  python scripts/build_dora_workbook.py
"""

from __future__ import annotations

import csv
import statistics
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generate_dora_data import PERIOD_DAYS, PERIOD_START, SEED  # noqa: E402
from xlsx_common import (  # noqa: E402
    AMBER_FILL,
    BODY,
    BOLD,
    BOX,
    GREEN_FILL,
    H2,
    INPUT,
    LINK,
    RED_FILL,
    header_row,
    paragraph,
    style_range,
    synthetic_banner,
    tidy_chart,
    title,
    widths,
)

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "dora_deployments.csv"
OUT = ROOT / "reports" / "dora-analysis.xlsx"
PHASES = ["1-Manual", "2-CI", "3-GitOps"]
SOURCE = (
    "DORA (Google Cloud), Accelerate State of DevOps Report 2024 - software delivery "
    "performance clusters. https://dora.dev/research/2024/dora-report/"
)
NAMES = {  # named range -> Raw_Data column
    "DeploymentId": "A",
    "DeployDate": "C",
    "Phase": "D",
    "LeadTimeHours": "K",
    "ChangeFailed": "L",
    "FailureCategory": "M",
    "RecoveryMinutes": "P",
    "Remediation": "Q",
    "ToilMinutes": "R",
}


def load_rows() -> list[dict]:
    rows = []
    for r in csv.DictReader(DATA.open(encoding="utf-8")):
        rows.append(
            {
                **r,
                "deployed_at": datetime.strptime(r["deployed_at_utc"], "%Y-%m-%dT%H:%M:%SZ"),
                "first_commit_at": datetime.strptime(
                    r["first_commit_at_utc"], "%Y-%m-%dT%H:%M:%SZ"
                ),
                "restored_at": datetime.strptime(r["restored_at_utc"], "%Y-%m-%dT%H:%M:%SZ")
                if r["restored_at_utc"]
                else None,
                "date": date.fromisoformat(r["deploy_date"]),
                "lead": float(r["lead_time_hours"]),
                "failed": r["change_failed"] == "TRUE",
                "recovery": int(r["recovery_minutes"]) if r["recovery_minutes"] else None,
                "toil": int(r["manual_toil_minutes"]),
                "pr": int(r["pr_number"]),
            }
        )
    return rows


def classify_df(per_day: float) -> str:
    if per_day >= 1:
        return "Elite"
    if per_day >= 1 / 7:
        return "High"
    return "Medium" if per_day >= 1 / 30 else "Low"


def main() -> None:
    rows = load_rows()
    n = len(rows)
    last = n + 1
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    # ------------------------------------------------------------- Raw_Data
    raw = wb.create_sheet("Raw_Data")
    heads = [
        "deployment_id",
        "deployed_at_utc",
        "deploy_date",
        "phase",
        "deploy_method",
        "service",
        "environment",
        "commit_sha",
        "pr_number",
        "first_commit_at_utc",
        "lead_time_hours",
        "change_failed",
        "failure_category",
        "incident_id",
        "restored_at_utc",
        "recovery_minutes",
        "remediation",
        "manual_toil_minutes",
    ]
    header_row(raw, 1, heads)
    for i, r in enumerate(rows, start=2):
        vals = [
            r["deployment_id"],
            r["deployed_at"],
            r["date"],
            r["phase"],
            r["deploy_method"],
            r["service"],
            r["environment"],
            r["commit_sha"],
            r["pr"],
            r["first_commit_at"],
            r["lead"],
            r["failed"],
            r["failure_category"] or None,
            r["incident_id"] or None,
            r["restored_at"],
            r["recovery"],
            r["remediation"] or None,
            r["toil"],
        ]
        for j, v in enumerate(vals, start=1):
            c = raw.cell(row=i, column=j, value=v)
            c.font = INPUT
    for col, fmt in (
        ("B", "yyyy-mm-dd hh:mm"),
        ("C", "yyyy-mm-dd"),
        ("J", "yyyy-mm-dd hh:mm"),
        ("K", "0.00"),
        ("O", "yyyy-mm-dd hh:mm"),
    ):
        for c in raw[f"{col}2:{col}{last}"]:
            c[0].number_format = fmt
    raw.freeze_panes = "B2"
    raw.auto_filter.ref = f"A1:R{last}"
    widths(raw, [11, 17, 11, 10, 24, 15, 20, 13, 9, 17, 10, 10, 22, 10, 17, 10, 24, 10])
    raw.conditional_formatting.add(f"A2:R{last}", FormulaRule(formula=["$L2=TRUE"], fill=RED_FILL))
    for name, col in NAMES.items():
        wb.defined_names[name] = DefinedName(name, attr_text=f"Raw_Data!${col}$2:${col}${last}")

    # --------------------------------------------------------------- Metrics
    m = wb.create_sheet("Metrics")
    title(
        m,
        "DORA metrics - calculated from Raw_Data",
        "Overall 90-day window and per 30-day phase. All cells are formulas.",
    )
    synthetic_banner(
        m,
        "A3",
        7,
        "SYNTHETIC dataset - simulated for the assignment, not real HealthQueue production data.",
    )
    header_row(
        m,
        5,
        [
            "Metric",
            "Unit",
            "Overall (90 d)",
            PHASES[0],
            PHASES[1],
            PHASES[2],
            "How it is calculated",
        ],
    )
    m["C4"] = "Period start"
    m["D4"] = PERIOD_START
    m["E4"] = "Period end"
    m["F4"] = PERIOD_START + timedelta(days=PERIOD_DAYS - 1)
    for c in ("D4", "F4"):
        m[c].number_format = "yyyy-mm-dd"
        m[c].font = INPUT
    metric_rows = [
        (
            6,
            "Days in window",
            "days",
            PERIOD_DAYS,
            [30, 30, 30],
            "Input: 90-day window = three 30-day phases.",
        ),
        (
            7,
            "Deployments",
            "count",
            "=COUNTA(DeploymentId)",
            "=COUNTIF(Phase,{p})",
            "COUNTA / COUNTIF over Raw_Data.",
        ),
        (
            8,
            "Deployment frequency",
            "deploys/day",
            "=C7/C6",
            "={c}7/{c}6",
            "Deployments / days in window.",
        ),
        (9, "Deployment frequency", "deploys/week", "=C8*7", "={c}8*7", "Deploys per day x 7."),
        (
            10,
            "Change lead time (median)",
            "hours",
            "=MEDIAN(LeadTimeHours)",
            "{{=MEDIAN(IF(Phase={p},LeadTimeHours))}}",
            "Median of (deployed_at - first_commit_at). Median resists outliers.",
        ),
        (
            11,
            "Change lead time (90th percentile)",
            "hours",
            "=PERCENTILE(LeadTimeHours,0.9)",
            "{{=PERCENTILE(IF(Phase={p},LeadTimeHours),0.9)}}",
            "Slowest 10% of changes take at least this long.",
        ),
        (
            12,
            "Failed deployments",
            "count",
            "=COUNTIF(ChangeFailed,TRUE)",
            "=COUNTIFS(Phase,{p},ChangeFailed,TRUE)",
            "Deployments that caused a production failure needing rollback/hotfix.",
        ),
        (
            13,
            "Change failure rate",
            "%",
            "=C12/C7",
            "={c}12/{c}7",
            "Failed deployments / deployments.",
        ),
        (
            14,
            "Failed deployment recovery time (median)",
            "minutes",
            "{=MEDIAN(IF(ChangeFailed=TRUE,RecoveryMinutes))}",
            "{{=MEDIAN(IF((Phase={p})*(ChangeFailed=TRUE),RecoveryMinutes))}}",
            "Median of (restored_at - deployed_at) for failed deployments.",
        ),
        (
            15,
            "Failed deployment recovery time (mean)",
            "minutes",
            "=AVERAGEIFS(RecoveryMinutes,ChangeFailed,TRUE)",
            "=AVERAGEIFS(RecoveryMinutes,Phase,{p},ChangeFailed,TRUE)",
            "Mean is shown to expose long-tail incidents.",
        ),
        (
            16,
            "Manual deployment toil",
            "hours",
            "=SUM(ToilMinutes)/60",
            "=SUMIF(Phase,{p},ToilMinutes)/60",
            "Hands-on minutes per deployment (build, deploy, smoke test, rollback), summed.",
        ),
        (
            17,
            "Manual toil per deployment",
            "minutes",
            "=C16*60/C7",
            "={c}16*60/{c}7",
            "Toil / deployments.",
        ),
    ]
    fmts = {
        6: "0",
        7: "0",
        8: "0.00",
        9: "0.0",
        10: "0.0",
        11: "0.0",
        12: "0",
        13: "0.0%",
        14: "0",
        15: "0",
        16: "0.0",
        17: "0.0",
    }
    for r, name, unit, overall, phase_f, how in metric_rows:
        m[f"A{r}"], m[f"B{r}"], m[f"G{r}"] = name, unit, how
        if isinstance(overall, str) and overall.startswith("{"):
            m[f"C{r}"] = ArrayFormula(f"C{r}", overall[1:-1])
        else:
            m[f"C{r}"] = overall
        for k, col in enumerate("DEF"):
            if isinstance(phase_f, list):
                m[f"{col}{r}"] = phase_f[k]
                continue
            f = phase_f.format(p=f"{col}$5", c=col)
            if f.startswith("{"):
                m[f"{col}{r}"] = ArrayFormula(f"{col}{r}", f[1:-1])
            else:
                m[f"{col}{r}"] = f
        for col in "CDEF":
            m[f"{col}{r}"].number_format = fmts[r]
    style_range(m, "A6:G17", BODY)
    style_range(m, "A6:A17", BOLD)
    style_range(m, "G6:G17", BODY, wrap=True)
    for col in "CDEF":
        m[f"{col}6"].font = INPUT
        for r in range(6, 18):
            m[f"{col}{r}"].number_format = fmts[r]
    widths(m, [40, 13, 15, 12, 12, 12, 60])
    paragraph(
        m,
        "A19",
        "Why the median? Lead and recovery times are skewed (a few changes wait for days). "
        "The median describes the typical change; the 90th percentile and mean show the tail. Small "
        "samples (e.g. only a handful of failures per phase) make the per-phase failure rate and "
        "recovery time noisy - treat single-phase differences of a few points with caution.",
        7,
        48,
    )

    # ---------------------------------------------------------- Weekly_Trend
    wk = wb.create_sheet("Weekly_Trend")
    title(wk, "Weekly trend", "13 weeks (the last week has 6 days). Formulas over Raw_Data.")
    synthetic_banner(wk, "A3", 9, "SYNTHETIC dataset.")
    header_row(
        wk,
        5,
        [
            "Week",
            "Start",
            "End",
            "Deployments",
            "Deploys/day",
            "Failed",
            "Change failure rate",
            "Median lead time (h)",
            "Median recovery (min)",
            "Toil (h)",
        ],
    )
    weeks = (PERIOD_DAYS + 6) // 7
    for w in range(1, weeks + 1):
        r = 5 + w
        wk[f"A{r}"] = w
        wk[f"B{r}"] = f"=Metrics!$D$4+7*(A{r}-1)"
        wk[f"C{r}"] = f"=MIN(B{r}+6,Metrics!$F$4)"
        rng = f'DeployDate,">="&B{r},DeployDate,"<="&C{r}'
        wk[f"D{r}"] = f"=COUNTIFS({rng})"
        wk[f"E{r}"] = f"=D{r}/(C{r}-B{r}+1)"
        wk[f"F{r}"] = f"=COUNTIFS({rng},ChangeFailed,TRUE)"
        wk[f"G{r}"] = f'=IF(D{r}=0,"",F{r}/D{r})'
        wk[f"H{r}"] = ArrayFormula(
            f"H{r}", f'=IFERROR(MEDIAN(IF((DeployDate>=B{r})*(DeployDate<=C{r}),LeadTimeHours)),"")'
        )
        wk[f"I{r}"] = ArrayFormula(
            f"I{r}",
            f'=IF(F{r}=0,"",MEDIAN(IF((DeployDate>=B{r})*(DeployDate<=C{r})*(ChangeFailed=TRUE),RecoveryMinutes)))',
        )
        wk[f"J{r}"] = f"=SUMIFS(ToilMinutes,{rng})/60"
    lw = 5 + weeks
    style_range(wk, f"A6:J{lw}", BODY)
    for col, fmt in (
        ("B", "yyyy-mm-dd"),
        ("C", "yyyy-mm-dd"),
        ("E", "0.00"),
        ("G", "0%"),
        ("H", "0.0"),
        ("I", "0"),
        ("J", "0.0"),
    ):
        for c in wk[f"{col}6:{col}{lw}"]:
            c[0].number_format = fmt
    widths(wk, [7, 12, 12, 12, 11, 8, 12, 14, 14, 9])

    def add_chart(chart, anchor, cols, cat_col="A"):
        for col in cols:
            chart.add_data(
                Reference(wk, range_string=f"Weekly_Trend!{col}5:{col}{lw}"), titles_from_data=True
            )
        chart.set_categories(Reference(wk, range_string=f"Weekly_Trend!{cat_col}6:{cat_col}{lw}"))
        tidy_chart(chart)
        chart.height, chart.width = 7.5, 17
        wk.add_chart(chart, anchor)

    c1 = BarChart()
    c1.title = "Deployments per week"
    add_chart(c1, "L5", ["D"])
    c1.legend = None
    c2 = LineChart()
    c2.title = "Median change lead time per week (hours)"
    add_chart(c2, "L21", ["H"])
    c2.legend = None
    c3 = LineChart()
    c3.title = "Change failure rate per week"
    c3.y_axis.number_format = "0%"
    add_chart(c3, "L37", ["G"])
    c3.legend = None
    c4 = BarChart()
    c4.title = "Manual deployment toil per week (hours)"
    add_chart(c4, "L53", ["J"])
    c4.legend = None

    # -------------------------------------------------------------- Benchmarks
    b = wb.create_sheet("Benchmarks")
    title(b, "Comparison with DORA 2024 performance clusters", SOURCE)
    header_row(
        b,
        4,
        [
            "Cluster (2024)",
            "Deployment frequency",
            "Change lead time",
            "Change failure rate",
            "Failed deployment recovery time",
        ],
    )
    clusters = [
        (
            "Elite",
            "On demand (multiple deploys per day)",
            "Less than one day",
            0.05,
            "Less than one hour",
        ),
        (
            "High",
            "Between once per day and once per week",
            "Between one day and one week",
            0.20,
            "Less than one day",
        ),
        (
            "Medium",
            "Between once per week and once per month",
            "Between one week and one month",
            0.10,
            "Less than one day",
        ),
        (
            "Low",
            "Between once per month and once every six months",
            "Between one month and six months",
            0.40,
            "Between one week and one month",
        ),
    ]
    for i, row in enumerate(clusters, start=5):
        for j, v in enumerate(row, start=1):
            b.cell(row=i, column=j, value=v).font = INPUT
        b.cell(row=i, column=4).number_format = "0%"
    style_range(b, "A5:E8", INPUT, wrap=True)
    for i in range(5, 9):
        b[f"D{i}"].number_format = "0%"
        b.row_dimensions[i].height = 30
    paragraph(
        b,
        "A9",
        "Source note: values are the 2024 report's cluster descriptions (typical values, not hard "
        "cut-offs). The 2024 Medium cluster had a LOWER change failure rate (10%) than High (20%) - the "
        "report notes this, so CFR is compared value-by-value below rather than as a strict ranking. The "
        "2025 DORA report replaced these four tiers with seven team archetypes, so 2024 is the most recent "
        "edition with elite/high/medium/low benchmarks.",
        5,
        62,
    )

    b["A11"] = "Thresholds used to place a measured value in a cluster (our operationalisation)"
    b["A11"].font = H2
    header_row(b, 12, ["Metric", "Elite if", "High if", "Medium if", "Low otherwise / note"])
    th = [
        ("Deployment frequency (deploys/day)", 1, 1 / 7, 1 / 30, ">= 1 per 6 months"),
        ("Change lead time (median hours)", 24, 168, 720, "< 4,380 h (6 months)"),
        ("Change failure rate", 0.05, 0.20, 0.10, "<= 40% (note: Medium value < High value)"),
        ("Recovery time (median minutes)", 60, 1440, 1440, "<= 43,200 min (1 month)"),
    ]
    for i, row in enumerate(th, start=13):
        for j, v in enumerate(row, start=1):
            b.cell(row=i, column=j, value=v).font = INPUT if j > 1 else BOLD
    style_range(b, "A13:E16", INPUT)
    for c in ("B13", "C13", "D13"):
        b[c].number_format = "0.000"
    for c in ("B15", "C15", "D15"):
        b[c].number_format = "0%"
    b["A13"].font = b["A14"].font = b["A15"].font = b["A16"].font = BOLD
    b["F13"] = ">= threshold"
    b["F14"] = "< threshold"
    b["F15"] = "<= value"
    b["F16"] = "< threshold"

    b["A18"] = "HealthQueue (synthetic) vs. benchmarks"
    b["A18"].font = H2
    header_row(b, 19, ["Metric", "Overall (90 d)", "1-Manual", "2-CI", "3-GitOps"])
    comp = [
        (20, "Deployment frequency (deploys/day)", 8, "0.00"),
        (22, "Change lead time (median hours)", 10, "0.0"),
        (24, "Change failure rate", 13, "0.0%"),
        (26, "Recovery time (median minutes)", 14, "0"),
    ]
    for r, label, mrow, fmt in comp:
        b[f"A{r}"] = label
        b[f"A{r + 1}"] = "   -> cluster"
        for col, mcol in zip("BCDE", "CDEF", strict=True):
            b[f"{col}{r}"] = f"=Metrics!{mcol}{mrow}"
            b[f"{col}{r}"].number_format = fmt
            b[f"{col}{r}"].font = LINK
            v = f"{col}{r}"
            if mrow == 8:
                f = f'=IF({v}>=$B$13,"Elite",IF({v}>=$C$13,"High",IF({v}>=$D$13,"Medium","Low")))'
            elif mrow == 10:
                f = f'=IF({v}<$B$14,"Elite",IF({v}<$C$14,"High",IF({v}<$D$14,"Medium","Low")))'
            elif mrow == 13:
                f = (
                    f'=IF({v}<=$B$15,"Elite (<=5%)",IF({v}<=$D$15,"Medium value (<=10%)",'
                    f'IF({v}<=$C$15,"High value (<=20%)",IF({v}<=0.4,"Low value (<=40%)","Worse than Low"))))'
                )
            else:
                f = f'=IF({v}<$B$16,"Elite",IF({v}<$C$16,"High / Medium (< 1 day)",IF({v}<=43200,"Low","Worse than Low")))'
            b[f"{col}{r + 1}"] = f
    style_range(b, "A20:E27", BODY)
    for r, *_ in comp:
        for col in "BCDE":
            b[f"{col}{r}"].font = LINK
    b.conditional_formatting.add(
        "B21:E27", FormulaRule(formula=['LEFT(B21,5)="Elite"'], fill=GREEN_FILL)
    )
    b.conditional_formatting.add(
        "B21:E27",
        FormulaRule(formula=['OR(LEFT(B21,4)="High",LEFT(B21,6)="Medium")'], fill=AMBER_FILL),
    )
    b.conditional_formatting.add(
        "B21:E27", FormulaRule(formula=['OR(LEFT(B21,3)="Low",LEFT(B21,5)="Worse")'], fill=RED_FILL)
    )
    widths(b, [36, 30, 30, 26, 30, 14])
    paragraph(
        b,
        "A29",
        "Limitations: (1) the data is SYNTHETIC; (2) DORA clusters come from survey responses "
        "about many organisations, our numbers are computed from deployment records - the measurement "
        "methods differ; (3) one service and ~30 deployments per phase is a small sample; (4) a team "
        "should compare itself mainly with its own trend, not chase a label.",
        5,
        55,
    )

    # -------------------------------------------------------- Failure_Analysis
    fa = wb.create_sheet("Failure_Analysis")
    title(fa, "What fails, and what recovers fastest", "Drives the toil-reduction recommendations.")
    header_row(
        fa, 4, ["Failure category", "Failures", "Share of failures", "Median recovery (min)"]
    )
    cats = sorted({r["failure_category"] for r in rows if r["failed"]})
    for i, cat in enumerate(cats, start=5):
        fa[f"A{i}"] = cat
        fa[f"B{i}"] = f"=COUNTIF(FailureCategory,A{i})"
        fa[f"C{i}"] = f"=B{i}/Metrics!$C$12"
        fa[f"D{i}"] = ArrayFormula(f"D{i}", f"=MEDIAN(IF(FailureCategory=A{i},RecoveryMinutes))")
    lc = 4 + len(cats)
    style_range(fa, f"A5:D{lc}", BODY)
    for i in range(5, lc + 1):
        fa[f"C{i}"].number_format = "0%"
        fa[f"D{i}"].number_format = "0"
    r0 = lc + 3
    header_row(
        fa, r0, ["Remediation", "Times used", "Median recovery (min)", "Mean recovery (min)"]
    )
    rems = sorted({r["remediation"] for r in rows if r["failed"]})
    for i, rem in enumerate(rems, start=r0 + 1):
        fa[f"A{i}"] = rem
        fa[f"B{i}"] = f"=COUNTIF(Remediation,A{i})"
        fa[f"C{i}"] = ArrayFormula(f"C{i}", f"=MEDIAN(IF(Remediation=A{i},RecoveryMinutes))")
        fa[f"D{i}"] = f"=AVERAGEIFS(RecoveryMinutes,Remediation,A{i})"
    lr = r0 + len(rems)
    style_range(fa, f"A{r0 + 1}:D{lr}", BODY)
    for i in range(r0 + 1, lr + 1):
        fa[f"C{i}"].number_format = fa[f"D{i}"].number_format = "0"
    widths(fa, [30, 12, 20, 20])
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Median recovery time by remediation (minutes)"
    ch.add_data(Reference(fa, min_col=3, min_row=r0, max_row=lr), titles_from_data=True)
    ch.set_categories(Reference(fa, min_col=1, min_row=r0 + 1, max_row=lr))
    tidy_chart(ch)
    ch.legend = None
    ch.height, ch.width = 7, 16
    fa.add_chart(ch, "F4")

    # ------------------------------------------------------------------ Toil
    tl = wb.create_sheet("Toil_Reduction")
    title(
        tl,
        "Toil reduction - connecting the DORA results to automation",
        "Toil = manual, repetitive, automatable work that grows with deployments (Google SRE book, ch. 5).",
    )
    header_row(tl, 4, ["Measure", "1-Manual", "2-CI", "3-GitOps"])
    toil_rows = [
        ("Deployments in phase", 7, "0"),
        ("Manual toil (hours per 30 days)", 16, "0.0"),
        ("Manual toil per deployment (minutes)", 17, "0.0"),
        ("Median change lead time (hours)", 10, "0.0"),
        ("Median recovery time (minutes)", 14, "0"),
    ]
    for i, (label, mrow, fmt) in enumerate(toil_rows, start=5):
        tl[f"A{i}"] = label
        for col, mcol in zip("BCD", "DEF", strict=True):
            tl[f"{col}{i}"] = f"=Metrics!{mcol}{mrow}"
            tl[f"{col}{i}"].number_format = fmt
    tl["A10"] = "Toil per deployment reduced vs. phase 1"
    tl["C10"] = "=1-C7/$B$7"
    tl["D10"] = "=1-D7/$B$7"
    tl["A11"] = "Hours saved per 30 days at phase-3 volume if phase-1 process were still used"
    tl["D11"] = "=D5*($B$7-D7)/60"
    style_range(tl, "A5:D11", BODY)
    style_range(tl, "B5:D9", LINK)
    for i, (_, _, fmt) in enumerate(toil_rows, start=5):
        for col in "BCD":
            tl[f"{col}{i}"].number_format = fmt
    tl["C10"].number_format = tl["D10"].number_format = "0%"
    tl["D11"].number_format = "0.0"
    widths(tl, [62, 14, 14, 14, 70])

    tl["A13"] = "Toil inventory: what was automated, and what is left"
    tl["A13"].font = H2
    header_row(tl, 14, ["Manual task (phase 1)", "Automated by", "Status", "Effect on DORA metric"])
    inventory = [
        (
            "Build image on a laptop, push by hand",
            "CI pipeline (.github/workflows/ci.yml, release.yml)",
            "Done",
            "Lead time: no waiting for 'the person who builds'.",
        ),
        (
            "Manual vulnerability check / no check",
            "Trivy gate in CI + Dependabot PRs",
            "Done",
            "CFR: vulnerable images never ship.",
        ),
        (
            "SSH/kubectl deploy script run by hand",
            "Argo CD sync from Git (GitOps)",
            "Done",
            "Deployment frequency: deploy = merge a PR.",
        ),
        (
            "Manual rollback (find old version, redeploy)",
            "git revert -> Argo CD sync (tested: scripts/gitops_verify.sh)",
            "Done",
            "Recovery time: minutes instead of hours.",
        ),
        (
            "Manual smoke test after each deploy",
            "Readiness probes + SLO burn-rate alerts",
            "Partly",
            "CFR/recovery: failures detected automatically.",
        ),
        (
            "Approving each image-digest PR by hand",
            "Auto-merge digest PRs when signature + provenance verify",
            "Next",
            "Lead time: removes the remaining hand-off.",
        ),
        (
            "Watching dashboards during rollout",
            "Progressive delivery (Argo Rollouts canary + analysis on SLIs)",
            "Next",
            "CFR: bad versions stop at 10% of traffic.",
        ),
        (
            "Configuration mistakes reaching prod",
            "Policy-as-code tests on manifests (kubeconform + Kyverno/OPA)",
            "Partly",
            "CFR: biggest failure category (see Failure_Analysis).",
        ),
    ]
    for i, row in enumerate(inventory, start=15):
        for j, v in enumerate(row, start=1):
            tl.cell(row=i, column=j, value=v)
        tl.row_dimensions[i].height = 30
    style_range(tl, f"A15:D{14 + len(inventory)}", BODY, wrap=True)
    tl.conditional_formatting.add("C15:C22", FormulaRule(formula=['C15="Done"'], fill=GREEN_FILL))
    tl.conditional_formatting.add("C15:C22", FormulaRule(formula=['C15="Partly"'], fill=AMBER_FILL))
    tl.conditional_formatting.add("C15:C22", FormulaRule(formula=['C15="Next"'], fill=RED_FILL))
    widths(tl, [46, 52, 10, 46])
    tc = BarChart()
    tc.title = "Manual toil per deployment by phase (minutes)"
    tc.add_data(
        Reference(tl, min_col=2, max_col=4, min_row=7, max_row=7),
        from_rows=True,
        titles_from_data=False,
    )
    tc.set_categories(Reference(tl, min_col=2, max_col=4, min_row=4, max_row=4))
    tidy_chart(tc)
    tc.legend = None
    tc.height, tc.width = 7, 14
    tl.add_chart(tc, "F4")

    # ----------------------------------------------------------- Definitions
    df = wb.create_sheet("Definitions")
    title(df, "Definitions, units, windows and formulas")
    header_row(
        df,
        3,
        ["Metric", "DORA definition (2024)", "How this workbook measures it", "Unit", "Window"],
    )
    defs = [
        (
            "Deployment frequency",
            "How often application changes are deployed to production.",
            "Count of production deployments / days in window.",
            "deploys per day (and per week)",
            "90 d and 30-d phases",
        ),
        (
            "Change lead time",
            "Time for a committed change to be successfully deployed to production.",
            "deployed_at_utc - first_commit_at_utc of the change; reported as median and p90.",
            "hours",
            "90 d and 30-d phases",
        ),
        (
            "Change failure rate",
            "Share of deployments that cause a failure in production requiring "
            "remediation (rollback, hotfix, roll-forward).",
            "Deployments with change_failed = TRUE / all deployments.",
            "%",
            "90 d and 30-d phases",
        ),
        (
            "Failed deployment recovery time",
            "Time to recover from a deployment that failed in production.",
            "restored_at_utc - deployed_at_utc for failed deployments; median and mean.",
            "minutes",
            "90 d and 30-d phases",
        ),
        (
            "Manual toil (supporting metric)",
            "Not a DORA metric. Google SRE definition of toil.",
            "manual_toil_minutes per deployment, summed.",
            "hours",
            "per phase",
        ),
    ]
    for i, row in enumerate(defs, start=4):
        for j, v in enumerate(row, start=1):
            df.cell(row=i, column=j, value=v)
        df.row_dimensions[i].height = 45
    style_range(df, "A4:E8", BODY, wrap=True)
    widths(df, [30, 50, 50, 18, 18])
    paragraph(
        df,
        "A10",
        f"Source of definitions and benchmarks: {SOURCE}. Synthetic data generated by "
        f"scripts/generate_dora_data.py with fixed seed {SEED}; the CSV is in data/dora_deployments.csv.",
        5,
        35,
    )

    # -------------------------------------------------------------- Summary
    title(
        summary,
        "HealthQueue DORA Metrics Analysis",
        f"90-day window {PERIOD_START:%d %b %Y} - {PERIOD_START + timedelta(days=PERIOD_DAYS - 1):%d %b %Y}. "
        "Built by scripts/build_dora_workbook.py.",
    )
    synthetic_banner(
        summary,
        "A3",
        6,
        "SYNTHETIC DATA - simulated deployment records for the assignment. Not real "
        "HealthQueue production performance.",
    )
    header_row(
        summary,
        5,
        [
            "DORA metric",
            "Overall (90 d)",
            "Cluster",
            "Last 30 d (3-GitOps)",
            "Cluster",
            "Trend vs. phase 1",
        ],
    )
    srows = [
        (
            "Deployment frequency (deploys/day)",
            20,
            "0.00",
            '=IF(Metrics!F8>Metrics!D8,"Improving","Worse")',
        ),
        (
            "Change lead time (median h)",
            22,
            "0.0",
            '=IF(Metrics!F10<Metrics!D10,"Improving","Worse")',
        ),
        ("Change failure rate", 24, "0.0%", '=IF(Metrics!F13<Metrics!D13,"Improving","Worse")'),
        ("Recovery time (median min)", 26, "0", '=IF(Metrics!F14<Metrics!D14,"Improving","Worse")'),
    ]
    for i, (label, brow, fmt, trend) in enumerate(srows, start=6):
        summary[f"A{i}"] = label
        summary[f"B{i}"] = f"=Benchmarks!B{brow}"
        summary[f"C{i}"] = f"=Benchmarks!B{brow + 1}"
        summary[f"D{i}"] = f"=Benchmarks!E{brow}"
        summary[f"E{i}"] = f"=Benchmarks!E{brow + 1}"
        summary[f"F{i}"] = trend
        summary[f"B{i}"].number_format = summary[f"D{i}"].number_format = fmt
    style_range(summary, "A6:F9", LINK)
    style_range(summary, "A6:A9", BOLD)
    for i, (_, _, fmt, _) in enumerate(srows, start=6):
        summary[f"B{i}"].number_format = summary[f"D{i}"].number_format = fmt
    summary.conditional_formatting.add(
        "C6:C9", FormulaRule(formula=['LEFT(C6,5)="Elite"'], fill=GREEN_FILL)
    )
    summary.conditional_formatting.add(
        "E6:E9", FormulaRule(formula=['LEFT(E6,5)="Elite"'], fill=GREEN_FILL)
    )
    summary.conditional_formatting.add(
        "C6:E9", FormulaRule(formula=['OR(LEFT(C6,4)="High",LEFT(C6,6)="Medium")'], fill=AMBER_FILL)
    )
    summary.conditional_formatting.add(
        "F6:F9", FormulaRule(formula=['F6="Improving"'], fill=GREEN_FILL)
    )
    summary.conditional_formatting.add("F6:F9", FormulaRule(formula=['F6="Worse"'], fill=RED_FILL))
    summary["A11"] = "Key findings"
    summary["A11"].font = H2
    cat_counts = {
        c: sum(1 for r in rows if r["failure_category"] == c)
        for c in sorted({r["failure_category"] for r in rows if r["failed"]})
    }
    top_n = max(cat_counts.values())
    top = [c for c, k in cat_counts.items() if k == top_n]
    n_failed = sum(cat_counts.values())
    top_failures_text = (
        f"4. Biggest failure sources: {' and '.join(top)} "
        f"({'each ' if len(top) > 1 else ''}{top_n} of {n_failed} failures"
        f"{', tied' if len(top) > 1 else ''}). Target them with policy-as-code checks on manifests/config and "
        "migration tests before merge (see Failure_Analysis)."
    )
    findings = [
        '="1. Throughput: deployments rose from "&TEXT(Metrics!D8,"0.00")&"/day (manual) to "&TEXT(Metrics!F8,"0.00")'
        '&"/day (GitOps); median lead time fell from "&TEXT(Metrics!D10,"0")&" h to "&TEXT(Metrics!F10,"0.0")&" h."',
        '="2. Stability: change failure rate went from "&TEXT(Metrics!D13,"0%")&" to "&TEXT(Metrics!F13,"0%")'
        '&" and median recovery from "&TEXT(Metrics!D14,"0")&" to "&TEXT(Metrics!F14,"0")&" minutes - speed did not cost stability."',
        '="3. Toil: hands-on minutes per deployment dropped "&TEXT(Toil_Reduction!D10,"0%")&" (from "&TEXT(Metrics!D17,"0")'
        '&" to "&TEXT(Metrics!F17,"0.0")&" min), saving ~"&TEXT(Toil_Reduction!D11,"0")&" engineer-hours per month at today\'s volume."',
        top_failures_text,
        "5. Caveat: phase 2 shows a lower failure rate than phase 3 - with ~30-60 deployments per phase that "
        "difference is within random noise; do not over-read single-phase changes.",
    ]
    for i, text in enumerate(findings, start=12):
        paragraph(summary, f"A{i}", text, 6, 34)
    summary["A18"] = "Recommendations (each tied to a metric)"
    summary["A18"].font = H2
    recs = [
        "Deployment frequency -> keep changes small; auto-merge digest-bump PRs once Cosign signature and SLSA "
        "provenance verify (removes the last manual hand-off).",
        "Lead time -> small PRs, trunk-based development and a PR review SLA; the p90 lead time (Metrics row 11) "
        "is several times the median, so the slowest changes - not the typical one - are the next target.",
        "Change failure rate -> policy-as-code on manifests/config and automated migration tests (the two "
        "largest failure categories); canary releases with Argo Rollouts gated on the SLO metrics.",
        "Recovery time -> make git-revert rollback the default remediation (fastest in Failure_Analysis); "
        "trigger it from the fast-burn SLO alert runbook.",
    ]
    for i, text in enumerate(recs, start=19):
        paragraph(summary, f"A{i}", text, 6, 32)
    summary["A24"] = "Independent calculation checks"
    summary["A24"].font = BOLD
    summary["C24"] = "=Checks!B3"
    summary["C24"].font = LINK
    summary.conditional_formatting.add(
        "C24", FormulaRule(formula=['C24="ALL CHECKS PASS"'], fill=GREEN_FILL)
    )
    summary.conditional_formatting.add(
        "C24", FormulaRule(formula=['C24<>"ALL CHECKS PASS"'], fill=RED_FILL)
    )
    widths(summary, [36, 16, 24, 20, 24, 18])

    # ---------------------------------------------------------------- Checks
    ck = wb.create_sheet("Checks")
    title(
        ck,
        "Independent calculation checks",
        "Python recomputes each metric from the CSV; Excel formulas must match.",
    )
    ck["A3"] = "Overall"
    ck["A3"].font = BOLD
    ck["B3"] = '=IF(COUNTIF(E6:E80,"FAIL")=0,"ALL CHECKS PASS","CHECK FAILURES")'
    ck["B3"].font = BOLD
    header_row(
        ck,
        5,
        [
            "Check",
            "Workbook formula value",
            "Independent Python value",
            "Abs. difference",
            "Result",
        ],
    )
    exp: dict[str, tuple[str, float]] = {}

    def add(label, ref, value):
        exp[label] = (ref, value)

    groups = {"Overall": rows, **{p: [r for r in rows if r["phase"] == p] for p in PHASES}}
    for (label, subset), col in zip(groups.items(), "CDEF", strict=True):
        days = PERIOD_DAYS if label == "Overall" else 30
        failed = [r for r in subset if r["failed"]]
        lead = sorted(r["lead"] for r in subset)
        add(f"{label}: deployments", f"Metrics!{col}7", len(subset))
        add(f"{label}: deploys/day", f"Metrics!{col}8", len(subset) / days)
        add(f"{label}: median lead time (h)", f"Metrics!{col}10", statistics.median(lead))
        # Excel PERCENTILE = linear interpolation between closest ranks (inclusive method)
        add(
            f"{label}: p90 lead time (h)",
            f"Metrics!{col}11",
            statistics.quantiles(lead, n=10, method="inclusive")[-1],
        )
        add(f"{label}: failed deployments", f"Metrics!{col}12", len(failed))
        add(f"{label}: change failure rate", f"Metrics!{col}13", len(failed) / len(subset))
        add(
            f"{label}: median recovery (min)",
            f"Metrics!{col}14",
            statistics.median(r["recovery"] for r in failed),
        )
        add(
            f"{label}: mean recovery (min)",
            f"Metrics!{col}15",
            statistics.fmean(r["recovery"] for r in failed),
        )
        add(f"{label}: toil hours", f"Metrics!{col}16", sum(r["toil"] for r in subset) / 60)
    # Lead time column must equal deployed - first commit (data integrity)
    add(
        "Data integrity: lead_time != deployed - first_commit (1 = violation)",
        "0",
        max(
            abs(r["lead"] - (r["deployed_at"] - r["first_commit_at"]).total_seconds() / 3600)
            for r in rows
        )
        > 0.01,
    )
    add(
        "Data integrity: recovery != restored - deployed (1 = violation)",
        "0",
        max(
            abs(r["recovery"] - (r["restored_at"] - r["deployed_at"]).total_seconds() / 60)
            for r in rows
            if r["failed"]
        )
        > 0.5,
    )
    w1 = [r for r in rows if r["date"] < PERIOD_START + timedelta(days=7)]
    add("Week 1 deployments", "Weekly_Trend!D6", len(w1))
    add("Sum of weekly deployments = total", "SUM(Weekly_Trend!D6:D18)", n)
    add(
        "Overall DF cluster is Elite (1=yes)",
        'IF(Benchmarks!B21="Elite",1,0)',
        1 if classify_df(n / PERIOD_DAYS) == "Elite" else 0,
    )
    for i, (label, (ref, value)) in enumerate(exp.items(), start=6):
        ck[f"A{i}"] = label
        ck[f"B{i}"] = f"={ref}"
        ck[f"C{i}"] = float(value)
        ck[f"C{i}"].font = INPUT
        ck[f"D{i}"] = f"=ABS(B{i}-C{i})"
        ck[f"E{i}"] = f'=IF(D{i}<=1E-6,"PASS","FAIL")'
    lk = 5 + len(exp)
    style_range(ck, f"A6:E{lk}", BODY, "0.0000")
    style_range(ck, f"C6:C{lk}", INPUT, "0.0000")
    ck.conditional_formatting.add(f"E6:E{lk}", FormulaRule(formula=['E6="PASS"'], fill=GREEN_FILL))
    ck.conditional_formatting.add(f"E6:E{lk}", FormulaRule(formula=['E6="FAIL"'], fill=RED_FILL))
    widths(ck, [52, 24, 24, 16, 10])
    for row in ck[f"A6:E{lk}"]:
        for c in row:
            c.border = BOX

    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(
        f"Wrote {OUT.relative_to(ROOT)} ({len(wb.sheetnames)} sheets, {n} deployments, {len(exp)} independent checks)"
    )


if __name__ == "__main__":
    main()
