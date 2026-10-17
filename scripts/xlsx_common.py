"""Shared openpyxl styling helpers for the HealthQueue report workbooks."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT = "Arial"
NAVY = "1F3864"
TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
SUBTITLE = Font(name=FONT, size=10, italic=True, color="595959")
H2 = Font(name=FONT, size=12, bold=True, color=NAVY)
HEADER = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
INPUT = Font(name=FONT, size=10, color="0000FF")  # blue = hard-coded input / assumption
LINK = Font(name=FONT, size=10, color="008000")  # green = pulled from another sheet
WARN = Font(name=FONT, size=10, bold=True, color="C00000")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
SYNTH_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def title(ws: Worksheet, text: str, subtitle: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = TITLE
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE
    ws.sheet_view.showGridLines = False


def synthetic_banner(ws: Worksheet, cell: str, width: int, text: str) -> None:
    ws[cell] = text
    ws[cell].font = WARN
    ws[cell].fill = SYNTH_FILL
    row = ws[cell].row
    col = ws[cell].column
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width - 1)


def header_row(ws: Worksheet, row: int, headers: list[str], start_col: int = 1) -> None:
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = HEADER
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BOX


def style_range(
    ws: Worksheet, rng: str, font: Font = BODY, fmt: str | None = None, wrap: bool = False
) -> None:
    for row in ws[rng]:
        for c in row:
            c.font = font
            c.border = BOX
            if fmt:
                c.number_format = fmt
            if wrap:
                c.alignment = WRAP


def widths(ws: Worksheet, values: dict[str, float] | list[float]) -> None:
    if isinstance(values, dict):
        for col, w in values.items():
            ws.column_dimensions[col].width = w
    else:
        for i, w in enumerate(values, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def paragraph(
    ws: Worksheet, cell: str, text: str, width_cols: int = 8, height: float = 45, font: Font = BODY
) -> None:
    ws[cell] = text
    ws[cell].font = font
    ws[cell].alignment = WRAP
    r, c = ws[cell].row, ws[cell].column
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + width_cols - 1)
    ws.row_dimensions[r].height = height


def tidy_chart(chart, legend_bottom: bool = True) -> None:
    """openpyxl >= 3.1 hides axes by default and Excel smooths/vary-colours
    series; make charts readable and consistent."""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    if hasattr(chart, "varyColors"):
        chart.varyColors = False
    for s in chart.series:
        if hasattr(s, "smooth"):
            s.smooth = False
    if chart.legend is not None and legend_bottom:
        chart.legend.position = "b"
        chart.legend.overlay = False
    # Axis titles collide with tick labels in Excel's default layout; the chart
    # title already says what is plotted.
    chart.x_axis.title = None
    chart.y_axis.title = None
    if chart.title is not None:
        chart.title.overlay = False  # keep the title above the plot area
    if getattr(chart, "dataLabels", None) is not None:
        chart.dataLabels.showSerName = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showLegendKey = False
