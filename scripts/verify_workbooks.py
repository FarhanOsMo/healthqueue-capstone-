"""Cross-platform workbook verification (no Excel required).

Checks, for each report workbook:
  * required sheets exist and metrics are formulas (not pasted numbers)
  * if the file carries cached values (saved by Excel or LibreOffice), there
    are no formula errors and the Checks sheet says "ALL CHECKS PASS"

Full recalculation needs a spreadsheet engine: on Windows with Excel run
scripts/verify_workbooks.ps1; elsewhere open the file in LibreOffice Calc
(free), press Ctrl+Shift+F9, and look at the Checks sheet.

Usage: python scripts/verify_workbooks.py
"""

from __future__ import annotations

import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REQUIRED = {
    "reports/slo-error-budget.xlsx": [
        "Dashboard",
        "SLI_Definitions",
        "Window_Targets",
        "Synthetic_Daily",
        "Error_Budget",
        "Observed_Local",
        "Error_Budget_Policy",
        "Explanation",
        "Checks",
    ],
    "reports/dora-analysis.xlsx": [
        "Summary",
        "Raw_Data",
        "Metrics",
        "Weekly_Trend",
        "Benchmarks",
        "Failure_Analysis",
        "Toil_Reduction",
        "Definitions",
        "Checks",
    ],
}
ERRORS = ("#DIV/0!", "#N/A", "#NAME?", "#REF!", "#VALUE!", "#NUM!", "#NULL!")


def main() -> int:
    warnings.simplefilter("ignore")
    ok = True
    for rel, sheets in REQUIRED.items():
        path = ROOT / rel
        if not path.exists():
            print(f"FAIL  {rel}: missing (run the build scripts)")
            ok = False
            continue
        wb = load_workbook(path)
        missing = [s for s in sheets if s not in wb.sheetnames]
        formulas = sum(
            1
            for ws in wb
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        )
        cached = load_workbook(path, data_only=True)
        errors = [
            f"{ws.title}!{c.coordinate}"
            for ws in cached
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value in ERRORS
        ]
        verdict = cached["Checks"]["B3"].value
        status = "PASS" if not missing and formulas > 100 and not errors else "FAIL"
        note = (
            f"checks='{verdict}'"
            if verdict is not None
            else "no cached values - open in Excel/LibreOffice to recalculate"
        )
        if verdict is not None and verdict != "ALL CHECKS PASS":
            status = "FAIL"
        ok &= status == "PASS"
        print(
            f"{status}  {rel}: sheets={len(wb.sheetnames)} formulas={formulas} "
            f"missing_sheets={missing or 'none'} cached_errors={len(errors)} {note}"
        )
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
