"""Build reports/slo-error-budget.xlsx from the synthetic SLI data and local load tests.

Every calculated number in the workbook is an Excel FORMULA (so it recalculates
if an input changes). The "Checks" sheet compares those formulas with values
computed independently here in Python; every row must show PASS.

Inputs:
  data/slo_daily_synthetic.csv                      (python scripts/generate_slo_data.py)
  reports/slo/local-load-test-v0.1.0-no-pool.json   (python scripts/measure_slis.py, optional)
  reports/slo/local-load-test-v0.1.1-pool.json      (optional)

Usage: python scripts/build_slo_workbook.py
"""

from __future__ import annotations

import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule

sys.path.insert(0, str(Path(__file__).resolve().parent))
from xlsx_common import (  # noqa: E402
    AMBER_FILL,
    BODY,
    BOLD,
    BOX,
    GREEN_FILL,
    H2,
    INPUT,
    LINK,
    RED_FILL,
    header_row,
    paragraph,
    style_range,
    synthetic_banner,
    tidy_chart,
    title,
    widths,
)

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "slo_daily_synthetic.csv"
LOAD_BEFORE = ROOT / "reports" / "slo" / "local-load-test-v0.1.0-no-pool.json"
LOAD_AFTER = ROOT / "reports" / "slo" / "local-load-test-v0.1.1-pool.json"
OUT = ROOT / "reports" / "slo-error-budget.xlsx"

WINDOW_DAYS = 30
SLO_AVAIL, SLO_LAT, SLO_SUB = 0.999, 0.99, 0.995
CAUTION, FREEZE = 0.50, 0.25
FIRST, LAST = 6, 35  # data rows on Synthetic_Daily
PCT = "0.00%"
PCT3 = "0.000%"
INT = "#,##0"


def status_formula(remaining_cell: str) -> str:
    return (
        f'=IF({remaining_cell}<0,"EXHAUSTED - feature freeze",'
        f'IF({remaining_cell}<Window_Targets!$B$21,"FREEZE risky releases",'
        f'IF({remaining_cell}<Window_Targets!$B$20,"CAUTION - extra review","HEALTHY")))'
    )


def main() -> None:
    rows = list(csv.DictReader(DATA.open(encoding="utf-8")))
    assert len(rows) == WINDOW_DAYS, "expected 30 daily rows"
    wb = Workbook()

    # ------------------------------------------------------------ Dashboard
    dash = wb.active
    dash.title = "Dashboard"

    # ------------------------------------------------------- Window_Targets
    wt = wb.create_sheet("Window_Targets")
    title(
        wt,
        "Measurement window, targets and allowed downtime",
        "Blue cells are design inputs (editable). Everything else is calculated.",
    )
    header_row(wt, 3, ["Parameter", "Value", "Explanation"])
    params = [
        ("Measurement window (days)", WINDOW_DAYS, "0", "Rolling 30-day window, evaluated daily."),
        ("Minutes in window", "=B4*24*60", INT, "30 x 24 x 60 = 43,200 minutes."),
        (
            "Availability SLO target",
            SLO_AVAIL,
            "0.0%",
            "Proposed design decision (not mandated by the assignment): 99.9% of eligible requests succeed.",
        ),
        (
            "Latency SLO target",
            SLO_LAT,
            "0.0%",
            "99% of eligible requests complete in under the threshold.",
        ),
        (
            "Latency threshold (ms)",
            300,
            "0",
            "Matches the le=0.3 histogram bucket exposed by the API.",
        ),
        (
            "Submission SLO target",
            SLO_SUB,
            "0.0%",
            "99.5% of VALID booking attempts get a correct answer (201 booked or 409 slot taken).",
        ),
        (
            "Allowed downtime in window (minutes)",
            "=B5*(1-B6)",
            "0.0",
            "Time-based budget: 43,200 x (1 - 0.999) = 43.2 minutes per 30 days.",
        ),
        ("Allowed downtime in window (hours)", "=B10/60", "0.00", "43.2 / 60 = 0.72 hours."),
    ]
    for i, (name, value, fmt, note) in enumerate(params, start=4):
        wt.cell(row=i, column=1, value=name)
        c = wt.cell(row=i, column=2, value=value)
        c.number_format = fmt
        wt.cell(row=i, column=3, value=note)
        c.font = INPUT if not str(value).startswith("=") else BODY
    style_range(wt, "A4:A11", BOLD)
    style_range(wt, "C4:C11", BODY, wrap=True)
    for r in range(4, 12):
        wt.cell(row=r, column=2).border = BOX
    for r in (4, 6, 7, 8, 9):
        wt.cell(row=r, column=2).font = INPUT

    wt["A13"] = "What different availability targets would allow (time-based, per window)"
    wt["A13"].font = H2
    header_row(
        wt,
        14,
        ["Availability target", "Allowed downtime / window (min)", "Allowed downtime / week (min)"],
    )
    for i, target in enumerate([0.99, 0.995, 0.999, 0.9995, 0.9999], start=15):
        wt.cell(row=i, column=1, value=target).number_format = "0.00%"
        wt.cell(row=i, column=1).font = INPUT
        wt.cell(row=i, column=2, value=f"=$B$5*(1-A{i})").number_format = "0.0"
        wt.cell(row=i, column=3, value=f"=7*24*60*(1-A{i})").number_format = "0.0"
    style_range(wt, "B15:C19", BODY, "0.0")
    style_range(wt, "A15:A19", INPUT, "0.00%")

    wt["A20"] = "Policy: CAUTION when budget remaining is below"
    wt["B20"] = CAUTION
    wt["A21"] = "Policy: FREEZE risky releases when remaining is below"
    wt["B21"] = FREEZE
    for r in (20, 21):
        wt[f"A{r}"].font = BOLD
        wt[f"B{r}"].font = INPUT
        wt[f"B{r}"].number_format = "0%"
    wt["C20"] = "See the Error_Budget_Policy sheet for what each state means."
    widths(wt, [52, 22, 90])

    # ------------------------------------------------------ SLI_Definitions
    sd = wb.create_sheet("SLI_Definitions")
    title(
        sd,
        "SLI / SLO definitions",
        "Each SLI = good events / eligible events, measured by the API's Prometheus metrics.",
    )
    cols = [
        "SLI",
        "User journey",
        "Numerator (good events)",
        "Denominator (eligible events)",
        "Excluded (not eligible)",
        "Window",
        "Success criterion",
        "SLO target",
        "PromQL (30-day)",
    ]
    header_row(sd, 4, cols)
    defs = [
        (
            "Availability (request-based)",
            "Any patient-facing API call",
            "Eligible /v1 requests answered with a non-5xx status",
            "All requests to /v1/* routes",
            "/healthz, /readyz, /metrics (operational endpoints); 4xx count as GOOD because the service worked correctly",
            "30-day rolling",
            "status < 500",
            "=Window_Targets!B6",
            'sum(rate(healthqueue_http_requests_total{route=~"/v1/.*",status!~"5.."}[30d])) / '
            'sum(rate(healthqueue_http_requests_total{route=~"/v1/.*"}[30d]))',
        ),
        (
            "Availability (time-based)",
            "Can a patient reach the API at all?",
            "Probe minutes in which the synthetic check succeeded",
            "All minutes in the window (43,200)",
            "Announced maintenance (none planned)",
            "30-day rolling",
            "probe returns 200 within 2 s",
            "=Window_Targets!B6",
            "1 - (bad_minutes / 43200)  (from blackbox probe: probe_success)",
        ),
        (
            "Latency",
            "Booking/lookup feels fast",
            "Eligible /v1 requests completed in < 300 ms (server-side)",
            "All requests to /v1/* routes",
            "Operational endpoints",
            "30-day rolling",
            "duration < 0.3 s",
            "=Window_Targets!B7",
            'sum(rate(healthqueue_http_request_duration_seconds_bucket{route=~"/v1/.*",le="0.3"}[30d])) / '
            'sum(rate(healthqueue_http_request_duration_seconds_count{route=~"/v1/.*"}[30d]))',
        ),
        (
            "Successful valid submissions",
            "A patient books an appointment",
            "Valid POST /v1/appointments answered correctly: 201 created or 409 slot taken",
            "Valid booking attempts = created + conflict + server error",
            "422 validation rejections (client mistakes, not service failures)",
            "30-day rolling",
            "outcome in {created, conflict}",
            "=Window_Targets!B9",
            'sum(increase(healthqueue_appointment_submissions_total{outcome=~"created|conflict"}[30d])) / '
            'sum(increase(healthqueue_appointment_submissions_total{outcome=~"created|conflict|error"}[30d]))',
        ),
    ]
    for i, d in enumerate(defs, start=5):
        for j, v in enumerate(d, start=1):
            sd.cell(row=i, column=j, value=v)
        sd.row_dimensions[i].height = 75
    style_range(sd, "A5:I8", BODY, wrap=True)
    style_range(sd, "H5:H8", LINK, "0.0%")
    widths(sd, [24, 22, 30, 26, 32, 12, 18, 10, 70])
    paragraph(
        sd,
        "A10",
        "Why request-based AND time-based? Request-based budgets count failed requests "
        "(e.g. 1,036 allowed errors out of ~1.04 M requests). Time-based budgets count minutes of "
        "downtime (43.2 min per 30 days). They answer different questions and are never mixed: "
        "a request failure is NOT a minute of downtime.",
        9,
        45,
    )

    # ------------------------------------------------------ Synthetic_Daily
    sy = wb.create_sheet("Synthetic_Daily")
    title(
        sy,
        "Daily SLI data - September 2026",
        "Generated by scripts/generate_slo_data.py (fixed seed).",
    )
    synthetic_banner(
        sy,
        "A3",
        15,
        "SYNTHETIC DATA - simulated for this assignment, NOT real HealthQueue production measurements.",
    )
    heads = [
        "Date",
        "Eligible requests",
        "5xx errors",
        "Slow requests (>=300 ms)",
        "Valid submissions",
        "Failed submissions",
        "Probe minutes",
        "Bad minutes",
        "Incident note",
        "Daily availability SLI",
        "Daily latency SLI",
        "Daily submission SLI",
        "Cumulative availability budget used",
        "Cumulative downtime budget used",
        "Availability target",
    ]
    header_row(sy, 5, heads)
    for i, r in enumerate(rows, start=FIRST):
        sy.cell(row=i, column=1, value=date.fromisoformat(r["date"])).number_format = "yyyy-mm-dd"
        for j, k in enumerate(
            [
                "eligible_requests",
                "errors_5xx",
                "slow_requests_ge_300ms",
                "valid_submissions",
                "failed_submissions",
                "probe_minutes",
                "bad_minutes",
            ],
            start=2,
        ):
            sy.cell(row=i, column=j, value=int(r[k])).number_format = INT
        sy.cell(row=i, column=9, value=r["incident_note"] or None)
        sy.cell(row=i, column=10, value=f"=1-C{i}/B{i}").number_format = PCT3
        sy.cell(row=i, column=11, value=f"=1-D{i}/B{i}").number_format = PCT
        sy.cell(row=i, column=12, value=f"=1-F{i}/E{i}").number_format = PCT3
        sy.cell(
            row=i, column=13, value=f"=SUM($C${FIRST}:C{i})/Error_Budget!$F$6"
        ).number_format = "0.0%"
        sy.cell(
            row=i, column=14, value=f"=SUM($H${FIRST}:H{i})/Error_Budget!$F$9"
        ).number_format = "0.0%"
        sy.cell(row=i, column=15, value="=Window_Targets!$B$6").number_format = "0.0%"
    style_range(sy, f"A{FIRST}:O{LAST}", BODY)
    for row in sy[f"B{FIRST}:H{LAST}"]:
        for c in row:
            c.font = INPUT
            c.number_format = INT
    sy[f"A{LAST + 1}"] = "Total"
    sy[f"A{LAST + 1}"].font = BOLD
    for col in "BCDEFGH":
        sy[f"{col}{LAST + 1}"] = f"=SUM({col}{FIRST}:{col}{LAST})"
        sy[f"{col}{LAST + 1}"].font = BOLD
        sy[f"{col}{LAST + 1}"].number_format = INT
    widths(sy, [12, 12, 10, 14, 12, 12, 10, 10, 46, 12, 12, 12, 14, 14, 12])
    sy.freeze_panes = "B6"
    sy.conditional_formatting.add(
        f"J{FIRST}:J{LAST}",
        CellIsRule(operator="lessThan", formula=["Window_Targets!$B$6"], fill=RED_FILL),
    )
    sy.conditional_formatting.add(
        f"K{FIRST}:K{LAST}",
        CellIsRule(operator="lessThan", formula=["Window_Targets!$B$7"], fill=RED_FILL),
    )
    sy.conditional_formatting.add(
        f"L{FIRST}:L{LAST}",
        CellIsRule(operator="lessThan", formula=["Window_Targets!$B$9"], fill=RED_FILL),
    )
    sy.conditional_formatting.add(
        f"I{FIRST}:I{LAST}", FormulaRule(formula=[f"LEN(I{FIRST})>0"], fill=AMBER_FILL)
    )

    # --------------------------------------------------------- Error_Budget
    eb = wb.create_sheet("Error_Budget")
    title(
        eb,
        "Error budget: allowed, consumed, remaining",
        "Budget = (1 - SLO) x eligible events. Status thresholds come from Window_Targets.",
    )
    synthetic_banner(
        eb, "A3", 12, "Based on SYNTHETIC data (Synthetic_Daily). Not production measurements."
    )
    header_row(
        eb,
        5,
        [
            "SLI",
            "SLO target",
            "Eligible (window)",
            "Bad events",
            "Observed SLI",
            "Allowed bad (budget)",
            "Budget consumed",
            "Consumed %",
            "Budget remaining",
            "Remaining %",
            "Status",
            "Unit",
        ],
    )
    d = f"Synthetic_Daily!${{c}}${FIRST}:${{c}}${LAST}"
    specs = [
        (6, "Availability (request-based)", "Window_Targets!B6", "B", "C", "requests"),
        (7, "Latency (< 300 ms)", "Window_Targets!B7", "B", "D", "requests"),
        (8, "Successful valid submissions", "Window_Targets!B9", "E", "F", "submissions"),
        (9, "Availability (time-based)", "Window_Targets!B6", "G", "H", "minutes"),
    ]
    for r, name, slo, elig, bad, unit in specs:
        eb[f"A{r}"] = name
        eb[f"B{r}"] = f"={slo}"
        eb[f"C{r}"] = f"=SUM({d.format(c=elig)})"
        eb[f"D{r}"] = f"=SUM({d.format(c=bad)})"
        eb[f"E{r}"] = f"=1-D{r}/C{r}"
        eb[f"F{r}"] = f"=C{r}*(1-B{r})"
        eb[f"G{r}"] = f"=D{r}"
        eb[f"H{r}"] = f"=G{r}/F{r}"
        eb[f"I{r}"] = f"=F{r}-G{r}"
        eb[f"J{r}"] = f"=I{r}/F{r}"
        eb[f"K{r}"] = status_formula(f"J{r}")
        eb[f"L{r}"] = unit
    style_range(eb, "A6:L9", BODY)
    style_range(eb, "A6:A9", BOLD)
    style_range(eb, "B6:B9", LINK, "0.0%")
    style_range(eb, "C6:D9", BODY, INT)
    style_range(eb, "E6:E9", BODY, PCT3)
    style_range(eb, "F6:G9", BODY, "#,##0.0")
    style_range(eb, "I6:I9", BODY, "#,##0.0")
    style_range(eb, "H6:H9", BODY, "0.0%")
    style_range(eb, "J6:J9", BODY, "0.0%")
    for rule_fill, op, val in (
        (RED_FILL, "lessThan", "Window_Targets!$B$21"),
        (AMBER_FILL, "lessThan", "Window_Targets!$B$20"),
        (GREEN_FILL, "greaterThanOrEqual", "Window_Targets!$B$20"),
    ):
        eb.conditional_formatting.add(
            "J6:J9", CellIsRule(operator=op, formula=[val], fill=rule_fill, stopIfTrue=True)
        )
    eb.conditional_formatting.add(
        "K6:K9", FormulaRule(formula=['LEFT(K6,7)="HEALTHY"'], fill=GREEN_FILL)
    )
    eb.conditional_formatting.add(
        "K6:K9", FormulaRule(formula=['LEFT(K6,7)="CAUTION"'], fill=AMBER_FILL)
    )
    eb.conditional_formatting.add(
        "K6:K9",
        FormulaRule(formula=['OR(LEFT(K6,6)="FREEZE",LEFT(K6,9)="EXHAUSTED")'], fill=RED_FILL),
    )
    widths(eb, [30, 10, 16, 12, 13, 15, 15, 12, 16, 12, 28, 12])
    eb["A11"] = "How to read this"
    eb["A11"].font = H2
    paragraph(
        eb,
        "A12",
        "Example (availability, request-based): budget = eligible requests x (1 - 0.999). "
        "Consumed % = 5xx errors / budget. Remaining % = 1 - consumed %. The time-based row works the "
        "same way in MINUTES: 43,200 x 0.001 = 43.2 allowed bad minutes. Status: >=50% remaining = "
        "HEALTHY, 25-50% = CAUTION, 0-25% = FREEZE risky releases, <0 = EXHAUSTED.",
        12,
        60,
    )

    # -------------------------------------------------------- Observed_Local
    ob = wb.create_sheet("Observed_Local")
    title(
        ob,
        "Observed: local load tests against the real API on kind",
        "Source: scripts/measure_slis.py -> reports/slo/*.json. Counts are blue inputs; SLIs are formulas.",
    )
    synthetic_banner(
        ob,
        "A3",
        4,
        "OBSERVED on a laptop (3,000 synthetic requests, ~90 s). Real code, NOT production traffic, NOT a 30-day window.",
    )
    header_row(
        ob, 5, ["Measure", "v0.1.0 (connection per request)", "v0.1.1 (connection pool)", "Notes"]
    )
    before = json.loads(LOAD_BEFORE.read_text()) if LOAD_BEFORE.exists() else None
    after = json.loads(LOAD_AFTER.read_text()) if LOAD_AFTER.exists() else None
    items = [
        ("Eligible /v1 requests (server)", "eligible_requests"),
        ("5xx errors (server)", "errors_5xx"),
        ("Requests <= 300 ms (server histogram)", "fast_requests_le_0_3s"),
        ("Timed requests (server histogram)", "timed_requests"),
        ("Submissions created", "submissions_created"),
        ("Submissions conflict (409)", "submissions_conflict"),
        ("Submissions error (5xx)", "submissions_error"),
        ("Submissions rejected (422, not eligible)", "submissions_rejected_validation"),
    ]
    for i, (label, key) in enumerate(items, start=6):
        ob.cell(row=i, column=1, value=label)
        for col, rep in ((2, before), (3, after)):
            v = rep["server_side_prometheus_delta"][key] if rep else None
            ob.cell(row=i, column=col, value=v).font = INPUT
    for i, (label, key) in enumerate(
        [
            ("Client p50 latency (ms)", "p50_ms"),
            ("Client p95 latency (ms)", "p95_ms"),
            ("Client p99 latency (ms)", "p99_ms"),
        ],
        start=14,
    ):
        ob.cell(row=i, column=1, value=label)
        for col, rep in ((2, before), (3, after)):
            ob.cell(row=i, column=col, value=rep["client_side"][key] if rep else None).font = INPUT
    sli_rows = [
        (18, "Availability SLI", "=1-{c}7/{c}6", "Window_Targets!$B$6"),
        (19, "Latency SLI (<= 300 ms)", "={c}8/{c}9", "Window_Targets!$B$7"),
        (20, "Submission SLI", "=({c}10+{c}11)/({c}10+{c}11+{c}12)", "Window_Targets!$B$9"),
    ]
    header_row(ob, 17, ["SLI (formula)", "v0.1.0", "v0.1.1", "Meets SLO?"])
    for r, label, f, target in sli_rows:
        ob[f"A{r}"] = label
        ob[f"B{r}"] = f.format(c="B")
        ob[f"C{r}"] = f.format(c="C")
        ob[f"D{r}"] = (
            f'="v0.1.0: "&IF(B{r}>={target},"yes","NO")&"   v0.1.1: "&IF(C{r}>={target},"yes","NO")'
        )
    style_range(ob, "A6:D20", BODY)
    style_range(ob, "B18:C20", BODY, "0.00%")
    style_range(ob, "B6:C16", INPUT, "#,##0.0")
    for r, *_, target in sli_rows:
        ob.conditional_formatting.add(
            f"B{r}:C{r}", CellIsRule(operator="lessThan", formula=[target], fill=RED_FILL)
        )
        ob.conditional_formatting.add(
            f"B{r}:C{r}",
            CellIsRule(operator="greaterThanOrEqual", formula=[target], fill=GREEN_FILL),
        )
    ob["D6"] = "Prometheus counter deltas summed across both API pods."
    ob["D14"] = "Measured by the load generator (includes network/NodePort)."
    widths(ob, [40, 30, 28, 60])
    paragraph(
        ob,
        "A22",
        "Finding: the first load test showed only ~85% of requests under 300 ms, so the latency "
        "SLO would have been violated. Root cause: a new PostgreSQL connection per request. Fix: a "
        "connection pool (v0.1.1), shipped through the full pipeline (Trivy, SBOM, Cosign) and deployed by "
        "a GitOps commit. Re-measured: 100% under 300 ms, p50 down from ~210 ms to ~64 ms. "
        "This is the observability loop working: measure -> find -> fix -> verify.",
        4,
        75,
    )

    # ---------------------------------------------------- Error_Budget_Policy
    po = wb.create_sheet("Error_Budget_Policy")
    title(po, "Error budget policy", "Agreed between product and engineering; reviewed monthly.")
    header_row(po, 4, ["Budget remaining", "State", "What the team does"])
    policy = [
        (
            ">= 50%",
            "HEALTHY",
            "Ship normally. Planned risky work (migrations, dependency majors) is allowed.",
        ),
        (
            "25% - 50%",
            "CAUTION",
            "Risky changes need a second reviewer + rollback plan; deploy in business hours only; "
            "prioritise the top reliability ticket in the next sprint.",
        ),
        (
            "0% - 25%",
            "FREEZE risky releases",
            "Only low-risk changes, bug fixes and reliability work. Every deploy "
            "gets canary verification. Postmortem actions take priority over features.",
        ),
        (
            "< 0% (exhausted)",
            "EXHAUSTED - feature freeze",
            "Feature releases stop until the 30-day window recovers "
            "above 0%. Only fixes, security patches and reliability work ship. "
            "Escalate to CTO; incident review required.",
        ),
        (
            "Fast-burn alert",
            "PAGE",
            "A 14.4x burn (2% of budget in 1 h) pages on-call immediately - see runbook.",
        ),
    ]
    for i, row in enumerate(policy, start=5):
        for j, v in enumerate(row, start=1):
            po.cell(row=i, column=j, value=v)
        po.row_dimensions[i].height = 48
    style_range(po, "A5:C9", BODY, wrap=True)
    for r, fill in ((5, GREEN_FILL), (6, AMBER_FILL), (7, RED_FILL), (8, RED_FILL), (9, RED_FILL)):
        po[f"B{r}"].fill = fill
    widths(po, [20, 28, 100])
    paragraph(
        po,
        "A11",
        "Exceptions: security fixes and changes that IMPROVE reliability may always ship. "
        "The policy applies per SLO: if only latency is in FREEZE, performance-risky changes pause, "
        "unrelated UI copy changes can still ship.",
        3,
        40,
    )

    # ------------------------------------------------------------ Explanation
    ex = wb.create_sheet("Explanation")
    title(
        ex,
        "What the numbers mean",
        "Text below is generated from live formulas where it quotes numbers.",
    )
    lines = [
        (
            '="Availability (request-based): "&TEXT(Error_Budget!D6,"#,##0")&" failed requests out of "&TEXT(Error_Budget!C6,"#,##0")'
            '&" eligible. Budget allowed "&TEXT(Error_Budget!F6,"#,##0")&" failures, so "&TEXT(Error_Budget!H6,"0.0%")'
            '&" of the budget is used -> "&Error_Budget!K6&"."'
        ),
        (
            '="Availability (time-based): "&TEXT(Error_Budget!D9,"0")&" bad minutes out of an allowed "'
            '&TEXT(Error_Budget!F9,"0.0")&" minutes ("&TEXT(Error_Budget!H9,"0.0%")&" used) -> "&Error_Budget!K9&"."'
        ),
        (
            '="Latency: "&TEXT(Error_Budget!D7,"#,##0")&" slow requests vs. "&TEXT(Error_Budget!F7,"#,##0")'
            '&" allowed ("&TEXT(Error_Budget!H7,"0.0%")&" used) -> "&Error_Budget!K7&". Most came from INC-202 on 19 Sep."'
        ),
        (
            '="Submissions: "&TEXT(Error_Budget!D8,"#,##0")&" failed valid bookings vs. "&TEXT(Error_Budget!F8,"#,##0")'
            '&" allowed ("&TEXT(Error_Budget!H8,"0.0%")&" used) -> "&Error_Budget!K8&"."'
        ),
        "Interpretation: two incidents (INC-201 PostgreSQL failover, INC-203 bad config) consumed most of the "
        "availability budget. Under the policy the team should freeze risky releases, finish the INC-201 "
        "postmortem actions (connection pooling + DB failover testing) and add config validation to CI before "
        "resuming feature work.",
        "Why 99.9% and not 99.99%? 99.99% allows only 4.3 minutes of downtime per 30 days - unrealistic for a "
        "small startup team with one region and manual on-call. 99.9% (43.2 min) is a common starting point for "
        "a patient-facing booking API; tighten it once 3 months of real data exist.",
        "Limitation: all monthly figures are SYNTHETIC. The Observed_Local sheet is real measurement of the "
        "real code, but on a laptop for ~90 seconds - it shows the method works, not production reliability.",
    ]
    for i, text in enumerate(lines, start=4):
        paragraph(ex, f"A{i}", text, 10, 48)
    widths(ex, [14] * 10)

    # ------------------------------------------------------------- Dashboard
    title(
        dash,
        "HealthQueue SLO & Error Budget Dashboard",
        "30-day window: 1-30 September 2026. Built by scripts/build_slo_workbook.py.",
    )
    synthetic_banner(
        dash,
        "A3",
        7,
        "SYNTHETIC monthly data for the assignment - not real production measurements. "
        "See Observed_Local for real local load tests.",
    )
    header_row(
        dash, 5, ["SLI", "SLO", "Observed", "Budget used", "Budget remaining", "Status", "Unit"]
    )
    for i, r in enumerate(range(6, 10), start=6):
        dash[f"A{i}"] = f"=Error_Budget!A{r}"
        dash[f"B{i}"] = f"=Error_Budget!B{r}"
        dash[f"C{i}"] = f"=Error_Budget!E{r}"
        dash[f"D{i}"] = f"=Error_Budget!H{r}"
        dash[f"E{i}"] = f"=Error_Budget!J{r}"
        dash[f"F{i}"] = f"=Error_Budget!K{r}"
        dash[f"G{i}"] = f"=Error_Budget!L{r}"
    style_range(dash, "A6:G9", LINK)
    style_range(dash, "B6:B9", LINK, "0.0%")
    style_range(dash, "C6:C9", LINK, PCT3)
    style_range(dash, "D6:E9", LINK, "0.0%")
    for rule_fill, op, val in (
        (RED_FILL, "lessThan", "Window_Targets!$B$21"),
        (AMBER_FILL, "lessThan", "Window_Targets!$B$20"),
        (GREEN_FILL, "greaterThanOrEqual", "Window_Targets!$B$20"),
    ):
        dash.conditional_formatting.add(
            "E6:E9", CellIsRule(operator=op, formula=[val], fill=rule_fill, stopIfTrue=True)
        )
    dash.conditional_formatting.add(
        "F6:F9", FormulaRule(formula=['LEFT(F6,7)="HEALTHY"'], fill=GREEN_FILL)
    )
    dash.conditional_formatting.add(
        "F6:F9", FormulaRule(formula=['LEFT(F6,7)="CAUTION"'], fill=AMBER_FILL)
    )
    dash.conditional_formatting.add(
        "F6:F9",
        FormulaRule(formula=['OR(LEFT(F6,6)="FREEZE",LEFT(F6,9)="EXHAUSTED")'], fill=RED_FILL),
    )
    dash["A11"] = "Allowed downtime for 99.9% over 30 days (minutes)"
    dash["E11"] = "=Window_Targets!B10"
    dash["E11"].number_format = "0.0"
    dash["A12"] = "Independent calculation checks"
    dash["E12"] = "=Checks!B3"
    for c in ("A11", "A12"):
        dash[c].font = BOLD
    dash["E11"].font = dash["E12"].font = LINK
    dash.conditional_formatting.add(
        "E12", FormulaRule(formula=['E12="ALL CHECKS PASS"'], fill=GREEN_FILL)
    )
    dash.conditional_formatting.add(
        "E12", FormulaRule(formula=['E12<>"ALL CHECKS PASS"'], fill=RED_FILL)
    )
    widths(dash, [34, 9, 12, 13, 18, 28, 12])

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Error budget remaining by SLI"
    bar.y_axis.title = "Remaining"
    bar.y_axis.number_format = "0%"
    bar.add_data(Reference(dash, min_col=5, min_row=5, max_row=9), titles_from_data=True)
    bar.set_categories(Reference(dash, min_col=1, min_row=6, max_row=9))
    bar.legend = None
    from openpyxl.chart.label import DataLabelList

    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 1
    tidy_chart(bar)
    bar.height, bar.width = 7, 16
    dash.add_chart(bar, "A14")

    line = LineChart()
    line.title = "Cumulative budget consumed through the month"
    line.y_axis.title = "Budget used"
    line.y_axis.number_format = "0%"
    line.add_data(
        Reference(sy, min_col=13, max_col=14, min_row=5, max_row=LAST), titles_from_data=True
    )
    line.set_categories(Reference(sy, min_col=1, min_row=FIRST, max_row=LAST))
    line.x_axis.number_format = "dd"
    line.y_axis.scaling.min = 0
    tidy_chart(line)
    line.height, line.width = 7, 16
    dash.add_chart(line, "A29")

    daily = LineChart()
    daily.title = "Daily availability SLI vs 99.9% target"
    daily.y_axis.number_format = "0.00%"
    daily.add_data(Reference(sy, min_col=10, min_row=5, max_row=LAST), titles_from_data=True)
    daily.add_data(Reference(sy, min_col=15, min_row=5, max_row=LAST), titles_from_data=True)
    daily.set_categories(Reference(sy, min_col=1, min_row=FIRST, max_row=LAST))
    daily.x_axis.number_format = "dd"
    daily.y_axis.scaling.min = 0.985
    daily.y_axis.scaling.max = 1.0
    tidy_chart(daily)
    daily.series[1].graphicalProperties.line.dashStyle = "dash"
    daily.series[1].graphicalProperties.line.solidFill = "C00000"
    daily.height, daily.width = 7, 16
    dash.add_chart(daily, "A44")

    # ----------------------------------------------------------------- Checks
    ck = wb.create_sheet("Checks")
    title(ck, "Independent calculation checks")
    ck["A3"] = "Overall"
    ck["A3"].font = BOLD
    ck["B3"] = '=IF(COUNTIF(E6:E40,"FAIL")=0,"ALL CHECKS PASS","CHECK FAILURES")'
    ck["B3"].font = BOLD
    header_row(
        ck,
        5,
        [
            "Check",
            "Workbook formula value",
            "Independent Python value",
            "Abs. difference",
            "Result",
        ],
    )
    tot = {k: sum(int(r[k]) for r in rows) for k in rows[0] if k not in ("date", "incident_note")}
    exp = {
        "Allowed downtime minutes (99.9%, 30 d) = 43.2": (
            "Window_Targets!B10",
            30 * 24 * 60 * (1 - SLO_AVAIL),
        ),
        "Total eligible requests": ("Error_Budget!C6", tot["eligible_requests"]),
        "Total 5xx errors": ("Error_Budget!D6", tot["errors_5xx"]),
        "Availability SLI (request)": (
            "Error_Budget!E6",
            1 - tot["errors_5xx"] / tot["eligible_requests"],
        ),
        "Allowed failed requests": ("Error_Budget!F6", tot["eligible_requests"] * (1 - SLO_AVAIL)),
        "Availability budget consumed %": (
            "Error_Budget!H6",
            tot["errors_5xx"] / (tot["eligible_requests"] * (1 - SLO_AVAIL)),
        ),
        "Availability budget remaining %": (
            "Error_Budget!J6",
            1 - tot["errors_5xx"] / (tot["eligible_requests"] * (1 - SLO_AVAIL)),
        ),
        "Latency budget consumed %": (
            "Error_Budget!H7",
            tot["slow_requests_ge_300ms"] / (tot["eligible_requests"] * (1 - SLO_LAT)),
        ),
        "Submission budget consumed %": (
            "Error_Budget!H8",
            tot["failed_submissions"] / (tot["valid_submissions"] * (1 - SLO_SUB)),
        ),
        "Time-based minutes in window": ("Error_Budget!C9", tot["probe_minutes"]),
        "Downtime budget consumed %": (
            "Error_Budget!H9",
            tot["bad_minutes"] / (30 * 1440 * (1 - SLO_AVAIL)),
        ),
        "Cumulative availability budget at day 30": (
            f"Synthetic_Daily!M{LAST}",
            tot["errors_5xx"] / (tot["eligible_requests"] * (1 - SLO_AVAIL)),
        ),
    }
    if before and after:
        sb, sa = before["server_side_prometheus_delta"], after["server_side_prometheus_delta"]
        exp["Local latency SLI v0.1.0"] = (
            "Observed_Local!B19",
            sb["fast_requests_le_0_3s"] / sb["timed_requests"],
        )
        exp["Local latency SLI v0.1.1"] = (
            "Observed_Local!C19",
            sa["fast_requests_le_0_3s"] / sa["timed_requests"],
        )
    for i, (name, (ref, value)) in enumerate(exp.items(), start=6):
        ck[f"A{i}"] = name
        ck[f"B{i}"] = f"={ref}"
        ck[f"C{i}"] = value
        ck[f"C{i}"].font = INPUT
        ck[f"D{i}"] = f"=ABS(B{i}-C{i})"
        ck[f"E{i}"] = f'=IF(D{i}<=1E-9,"PASS","FAIL")'
    last = 5 + len(exp)
    style_range(ck, f"A6:E{last}", BODY, "0.000000")
    style_range(ck, f"C6:C{last}", INPUT, "0.000000")
    ck.conditional_formatting.add(
        f"E6:E{last}", FormulaRule(formula=['E6="PASS"'], fill=GREEN_FILL)
    )
    ck.conditional_formatting.add(f"E6:E{last}", FormulaRule(formula=['E6="FAIL"'], fill=RED_FILL))
    widths(ck, [46, 24, 24, 16, 10])

    wb.calculation.fullCalcOnLoad = True  # Excel/LibreOffice recalculate when opened
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(
        f"Wrote {OUT.relative_to(ROOT)} ({len(wb.sheetnames)} sheets, {len(exp)} independent checks)"
    )


if __name__ == "__main__":
    main()
